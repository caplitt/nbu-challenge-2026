"""Readers for NBU Excel files: Inflation Report chart data and the macro forecast table."""
from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

RAW = Path("data/raw/nbu")
IR_CHARTS = RAW / "Tables_Charts_2026-Q3.xlsx"
IR_FORECAST = RAW / "Macroeconomic_Forecast_table_2026_07.xlsx"

QLAB = re.compile(r"^(I|II|III|IV)\.(\d{2})$")
ROMAN = {"I": 1, "II": 2, "III": 3, "IV": 4}


def _num(v):
    if v is None:
        return np.nan
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", ".").strip())
    except ValueError:
        return np.nan


def read_ir_chart(sheet: str, path=IR_CHARTS, lang: str = "ua") -> pd.DataFrame:
    """Data behind one Inflation Report chart, indexed by Period (M or Q) or year.

    Layout of the NBU file (Tables_Charts_*.xlsx): one sheet per chart, row 2 = Ukrainian series names, row 3 =
    English names, data from row 4; the first column (or the second, when the first holds a year and the second a
    quarter label like 'IV.26') carries the period as a date, a quarter label or a two-digit year; blank labels
    continue the previous quarter. Numbers may be stored as text with a decimal comma."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
    data = rows[3:]
    label_col = 1 if any(isinstance(r[1], str) and QLAB.match(r[1].replace("І", "I").strip())
                         for r in data if len(r) > 1) else 0
    name_row = rows[1] if lang == "ua" else rows[2]
    names = []
    for v in name_row[label_col + 1:]:
        if v is None:
            break
        names.append(str(v).strip())
    index, values = [], []
    last_q = None
    for r in data:
        vals = [_num(v) for v in r[label_col + 1: label_col + 1 + len(names)]]
        lab = r[label_col] if label_col < len(r) else None
        if isinstance(lab, str):
            lab = lab.replace("І", "I").strip()   # Cyrillic І in some quarter labels
        if isinstance(lab, (dt.datetime, dt.date)):
            per = pd.Period(lab, "M")
        elif isinstance(lab, str) and QLAB.match(lab.strip()):
            q, y = QLAB.match(lab.strip()).groups()
            per = pd.Period(year=2000 + int(y), quarter=ROMAN[q], freq="Q")
        elif isinstance(lab, (str, int, float)) and re.fullmatch(r"\d{2}", str(lab).strip().split(".")[0]):
            per = 2000 + int(str(lab).strip().split(".")[0])
        elif lab is None and last_q is not None and not all(np.isnan(vals)):
            per = last_q + 1
        else:
            if all(np.isnan(vals)):
                continue
            per = lab
        if isinstance(per, pd.Period) and per.freqstr.startswith("Q"):
            last_q = per
        if all(np.isnan(vals)):
            continue
        index.append(per)
        values.append(vals)
    df = pd.DataFrame(values, index=index, columns=names)
    if len(df) and all(isinstance(i, pd.Period) for i in df.index):
        df.index = pd.PeriodIndex(df.index)
    return df.loc[:, ~df.columns.duplicated()]


def read_nbu_forecast(path=IR_FORECAST) -> pd.DataFrame:
    """NBU macro forecast table (sheet 'Базовий сценарій') in long form:
    indicator, period ('2025', '2026Q3', '2026', '2026_prev'), value."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = [list(r) for r in wb["Базовий сценарій"].iter_rows(values_only=True)]
    h = next(i for i, r in enumerate(rows) if "Показники" in [str(v).strip() for v in r if v is not None])
    years, flags, quarters = rows[h], rows[h + 1], rows[h + 2]
    cols = {}
    year = None
    for j in range(2, len(years)):
        if years[j] is not None:
            year = int(years[j])
        if year is None:
            continue
        if quarters[j] in ROMAN:
            cols[j] = f"{year}Q{ROMAN[quarters[j]]}"
        elif flags[j] == "прогноз 04.2026":
            cols[j] = f"{year}_prev"
        else:
            cols[j] = str(year)
    out = []
    for r in rows[h + 3:]:
        name = r[1]
        if not isinstance(name, str) or all(v is None for v in r[2:]):
            continue
        for j, lab in cols.items():
            v = _num(r[j]) if j < len(r) else np.nan
            if not np.isnan(v):
                out.append((name.strip(), lab, v))
    return pd.DataFrame(out, columns=["indicator", "period", "value"])


def nbu_forecast_series(indicator_prefix: str, df: pd.DataFrame | None = None) -> pd.Series:
    """Quarterly path of one indicator from the NBU forecast table (prefix match)."""
    df = read_nbu_forecast() if df is None else df
    sel = df[df["indicator"].str.startswith(indicator_prefix) & df["period"].str.contains("Q")]
    s = pd.Series(sel["value"].values, index=pd.PeriodIndex(sel["period"], freq="Q"))
    return s[~s.index.duplicated()]
